"""Write the solo-week spreadsheet from solo.json."""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC = sys.argv[1]
DST = sys.argv[2]
rows = json.load(open(SRC))

FONT = 'Arial'
H = Font(name=FONT, bold=True, color='FFFFFF')
B = Font(name=FONT)
BB = Font(name=FONT, bold=True)
LINK = Font(name=FONT, color='0563C1', underline='single')
MUTED = Font(name=FONT, color='666666', italic=True)
HEAD_FILL = PatternFill('solid', fgColor='1F2A44')
TOP_FILL = PatternFill('solid', fgColor='EAF2E3')
THIN = Side(style='thin', color='D0D0D0')
GRID = Border(bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')

wb = Workbook()

# ---- Sheet 1: the week at a glance ----
ws = wb.active
ws.title = 'Week plan'
ws['A1'] = 'Tokyo solo week, based in Shibuya'
ws['A1'].font = Font(name=FONT, bold=True, size=14)
ws['A2'] = 'Mon 26 – Fri 30 Oct 2026. Top pick first, then backups. Every place is open on that day by its listed closing days.'
ws['A2'].font = MUTED
head = ['Day', 'Meal', 'Pick', 'Place', 'Japanese name', 'Genre', 'Area', 'From Shibuya', 'Tabelog', 'Reviews',
        'Price', 'How to book', 'Why / note', 'Options open that day']
ws.append([])
ws.append(head)
hr = ws.max_row
for c in ws[hr]:
    c.font, c.fill, c.alignment = H, HEAD_FILL, Alignment(vertical='center', wrap_text=True)
picked = [r for r in rows if r['pick']]
first_data = hr + 1
for r in picked:
    ws.append([r['day'], r['meal'], r['pick'], r['name'], r['name_ja'], r['genre'], r['area'], r['distance'],
               r['tabelog'], r['reviews'], r['price'], r['booking'], r['note'], None])
    i = ws.max_row
    # Count of all options open that day, from the 'All options' sheet.
    ws.cell(i, 14, f"=COUNTIFS('All options'!$A:$A,A{i},'All options'!$B:$B,B{i})")
    if r['tabelog_url']:
        c = ws.cell(i, 4)
        c.hyperlink = r['tabelog_url']
        c.font = LINK
    for c in ws[i]:
        if c.column != 4 or not r['tabelog_url']:
            c.font = BB if r['pick'] == 'Top pick' and c.column in (3, 4) else B
        c.alignment = WRAP
        c.border = GRID
        if r['pick'] == 'Top pick':
            c.fill = TOP_FILL
    ws.cell(i, 9).number_format = '0.00'
    ws.cell(i, 10).number_format = '#,##0'
widths = [11, 8, 10, 30, 18, 22, 24, 22, 8, 9, 22, 40, 52, 11]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = ws.cell(first_data, 5)

# ---- Sheet 2: every option, filterable ----
wa = wb.create_sheet('All options')
cols = ['Day', 'Meal', 'Pick', 'Place', 'Japanese name', 'Genre', 'Area', 'From Shibuya', 'Tabelog', 'Reviews',
        'Awards', 'Price', 'How to book', 'Closed on', 'Note', 'Source', 'Tabelog link']
wa.append(cols)
for c in wa[1]:
    c.font, c.fill, c.alignment = H, HEAD_FILL, Alignment(vertical='center', wrap_text=True)
for r in rows:
    wa.append([r['day'], r['meal'], r['pick'], r['name'], r['name_ja'], r['genre'], r['area'], r['distance'],
               r['tabelog'], r['reviews'], r['awards'], r['price'], r['booking'], r['closed_days'], r['note'],
               r['source'], r['tabelog_url']])
    i = wa.max_row
    for c in wa[i]:
        c.font, c.alignment, c.border = B, WRAP, GRID
    if r['tabelog_url']:
        lc = wa.cell(i, 17)
        lc.hyperlink = r['tabelog_url']
        lc.font = LINK
    wa.cell(i, 9).number_format = '0.00'
    wa.cell(i, 10).number_format = '#,##0'
    if r['pick'] == 'Top pick':
        for c in wa[i]:
            c.fill = TOP_FILL
widths = [11, 8, 10, 30, 18, 22, 24, 22, 8, 9, 30, 22, 40, 14, 52, 24, 40]
for i, w in enumerate(widths, 1):
    wa.column_dimensions[get_column_letter(i)].width = w
wa.freeze_panes = 'E2'
wa.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{wa.max_row}'

# ---- Sheet 3: notes ----
wn = wb.create_sheet('Notes')
notes = [
    ('How to read this', ''),
    ('Pick', 'The concierge’s order for that meal. Blank means a good backup, not ranked.'),
    ('From Shibuya', 'Approximate, from the map: walk, one stop (Ebisu, Daikanyama, Nakameguro, Omotesandō), short cab of 10–15 min, or cross-town cab of 20–35 min. Not timed.'),
    ('Tabelog', 'Score and review count as seen on Tabelog during research (18–24 Oct 2026). 3.50+ is excellent; ramen and soba run lower, so 3.6+ there is very strong.'),
    ('Open', 'Each row is open that day by its listed weekly closing day. Irregular days off are not shown: check the Tabelog page or the shop’s own account the day before.'),
    ('Booking a week out', 'Top counters are full a week out. Solo seats come back as cancellations: check the Tabelog reservation page, or OMAKASE, morning and evening.'),
    ('Options open that day', 'A count of every option on the All options sheet for that day and meal.'),
    ('Sources', 'Tabelog only for quality; no Google or TripAdvisor ratings. Critics are named where a public link backs them.'),
]
for k, v in notes:
    wn.append([k, v])
wn['A1'].font = Font(name=FONT, bold=True, size=13)
for row in wn.iter_rows(min_row=2):
    row[0].font = BB
    row[1].font = B
    row[1].alignment = WRAP
wn.column_dimensions['A'].width = 24
wn.column_dimensions['B'].width = 100

wb.save(DST)
print('saved', DST, len(picked), 'picks', len(rows), 'options')
